"""Generate PDF, DOCX, PPT, and email drafts from intelligence bundle + health snapshot."""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

logger = logging.getLogger(__name__)


def _rag_rgb(rag: str) -> tuple[int, int, int]:
    if rag == "Green":
        return 0, 176, 80
    if rag == "Amber":
        return 255, 192, 0
    return 255, 0, 0


def write_executive_pdf(path: Path, title: str, markdown_body: str) -> Path:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    path.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    story: list[Any] = [Paragraph(escape(title), styles["Title"]), Spacer(1, 12)]
    for para in markdown_body.split("\n\n"):
        for line in para.split("\n"):
            t = line.strip() or " "
            story.append(Paragraph(escape(t)[:3500], styles["BodyText"]))
            story.append(Spacer(1, 4))
    doc = SimpleDocTemplate(str(path), pagesize=letter)
    doc.build(story)
    return path


def write_client_docx(path: Path, title: str, markdown_body: str) -> Path:
    from docx import Document
    from docx.shared import Pt

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title, 0)
    for para in markdown_body.split("\n\n"):
        p = doc.add_paragraph()
        for line in para.split("\n"):
            run = p.add_run(line.strip() + "\n")
            run.font.size = Pt(11)
    doc.save(str(path))
    return path


def write_intel_ppt(
    path: Path,
    project_name: str,
    health: dict[str, Any],
    bundle: dict[str, Any],
) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    rag = (health.get("rag") or {}).get("status", "Green")
    rgb = RGBColor(*_rag_rgb(str(rag)))

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = f"Intelligence report — {project_name}"
    slide.placeholders[1].text = datetime.now().strftime("%d %B %Y")

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Forecast"
    tf = slide.placeholders[1].text_frame
    fc = bundle.get("forecast") or {}
    tf.text = fc.get("headline", "")
    for line in fc.get("drivers") or []:
        tf.add_paragraph().text = str(line)

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Root causes"
    tf = slide.placeholders[1].text_frame
    first = True
    for c in (bundle.get("root_causes") or [])[:8]:
        t = f"[{c.get('severity')}] {c.get('statement')}"
        if first:
            tf.text = t
            first = False
        else:
            tf.add_paragraph().text = t

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Recommendations"
    tf = slide.placeholders[1].text_frame
    first = True
    for r in (bundle.get("recommendations") or [])[:8]:
        t = f"P{r.get('priority')} — {r.get('title')}"
        if first:
            tf.text = t
            first = False
        else:
            p = tf.add_paragraph()
            p.text = t
            p.level = 0

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "RAG"
    tf = slide.placeholders[1].text_frame
    p = tf.paragraphs[0]
    p.text = f"Status: {rag}"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = rgb

    prs.save(str(path))
    return path


def write_email_txt(path: Path, project_name: str, bundle: dict[str, Any]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    fc = bundle.get("forecast") or {}
    lines = [
        f"Subject: Project intelligence — {project_name}",
        "",
        "Dear team,",
        "",
        f"Forecast headline: {fc.get('headline', '')}",
        "",
        "Top recommendations:",
    ]
    for r in (bundle.get("recommendations") or [])[:6]:
        lines.append(f"- P{r.get('priority')}: {r.get('title')} — {r.get('detail')}")
    lines.extend(
        [
            "",
            "See attached/bundled PDF and DOCX for executive and client narratives.",
            "",
            "Regards,",
            "ProjectPilot PMO",
        ],
    )
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_markdown_bundle(out_dir: Path, bundle: dict[str, Any]) -> dict[str, Path]:
    """Write three markdown files; keys are stable API field names."""
    out_dir.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    mapping = (
        ("executive_summary.md", "executive_summary_markdown", "executive_summary_md"),
        ("pm_detailed_report.md", "pm_detailed_report_markdown", "pm_detailed_report_md"),
        ("client_report.md", "client_report_markdown", "client_report_md"),
    )
    for fname, bundle_key, out_key in mapping:
        p = out_dir / fname
        p.write_text(bundle.get(bundle_key) or "", encoding="utf-8")
        paths[out_key] = p
    return paths
