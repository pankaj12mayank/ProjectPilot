from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def _rag_color(rag: str) -> RGBColor:
    if rag == "Green":
        return RGBColor(0, 176, 80)
    if rag == "Amber":
        return RGBColor(255, 192, 0)
    return RGBColor(255, 0, 0)


def generate_ppt(
    metrics: tuple[float, float, float, int],
    chart_paths: tuple[Path, Path],
    rag: str,
    evm: tuple[float, float],
    output_dir: Path,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation()

    completion, total_sv, total_ev, risk_score = metrics
    spi, cpi = evm
    completion_chart, trend_chart = chart_paths

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Weekly Project Governance Report"
    slide.placeholders[1].text = datetime.today().strftime("%d %B %Y")

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Executive Summary"
    tf = slide.placeholders[1].text_frame
    tf.text = f"Overall Completion: {round(completion, 2)}%"
    tf.add_paragraph().text = f"Schedule Variance: {round(total_sv, 2)}"
    tf.add_paragraph().text = f"Effort Variance: {round(total_ev, 2)}"
    tf.add_paragraph().text = f"High Risks: {risk_score}"
    tf.add_paragraph().text = f"SPI: {round(spi, 2)} | CPI: {round(cpi, 2)}"

    p = tf.add_paragraph()
    p.text = f"Overall Health: {rag}"
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = _rag_color(rag)

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Completion Overview"
    slide.shapes.add_picture(str(completion_chart), Inches(1), Inches(1.5), width=Inches(6))

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Trend Overview"
    slide.shapes.add_picture(str(trend_chart), Inches(1), Inches(1.5), width=Inches(6))

    file_path = output_dir / f"Governance_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx"
    prs.save(str(file_path))
    return file_path
