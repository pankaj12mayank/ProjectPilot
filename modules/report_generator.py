from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from datetime import datetime
from config.settings import OUTPUT_FOLDER
import os

def get_color(rag):
    if rag == "Green":
        return RGBColor(0,176,80)
    elif rag == "Amber":
        return RGBColor(255,192,0)
    else:
        return RGBColor(255,0,0)

def generate_ppt(metrics, charts, rag, evm):

    os.makedirs(OUTPUT_FOLDER, exist_ok=True)
    prs = Presentation()

    completion, total_sv, total_ev, risk_score = metrics
    SPI, CPI = evm
    completion_chart, trend_chart = charts

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Weekly Project Governance Report"
    slide.placeholders[1].text = datetime.today().strftime("%d %B %Y")

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Executive Summary"
    tf = slide.placeholders[1].text_frame
    tf.text = f"Overall Completion: {round(completion,2)}%"
    tf.add_paragraph().text = f"Schedule Variance: {round(total_sv,2)}"
    tf.add_paragraph().text = f"Effort Variance: {round(total_ev,2)}"
    tf.add_paragraph().text = f"High Risks: {risk_score}"
    tf.add_paragraph().text = f"SPI: {round(SPI,2)} | CPI: {round(CPI,2)}"

    p = tf.add_paragraph()
    p.text = f"Overall Health: {rag}"
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = get_color(rag)

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Completion Overview"
    slide.shapes.add_picture(completion_chart, Inches(1), Inches(1.5), width=Inches(6))

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Trend Overview"
    slide.shapes.add_picture(trend_chart, Inches(1), Inches(1.5), width=Inches(6))

    file_path = f"{OUTPUT_FOLDER}/Governance_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx"
    prs.save(file_path)

    return file_path
